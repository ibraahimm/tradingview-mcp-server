#!/usr/bin/env python3
"""Canon comparison-view generator — regenerates .claude/outputs/tasi-rules-comparison.{md,xlsx}
from research/spec/rules.yaml, per the generated-Canon-view convention: a self-describing scope
banner + generator-proven completeness checks + the full executable definition (funnel predicates
with operator bound semantics, resolved thresholds marking @param (tunable) vs literal
(structural), null_action, enabled_when), never a params-only view.

    python research/spec/conformance/gen_canon_view.py [date] [commit_label]

Defaults: date = today, commit_label = `git rev-parse --short HEAD` (suffix "+dirty" when
rules.yaml has uncommitted changes). Run after ANY canon change. This is TOOLING, not a CI gate:
the artifacts are gitignored outputs; the generator asserts its own completeness checks and
exits non-zero if the view would be partial.
"""
from __future__ import annotations
import datetime as dt
import subprocess
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[3]


def default_commit() -> str:
    try:
        sha = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT,
                             capture_output=True, text=True, check=True).stdout.strip()
        dirty = subprocess.run(["git", "status", "--porcelain", "research/spec/rules.yaml"],
                               cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()
        return sha + ("+dirty" if dirty else "")
    except Exception:
        return "unknown"


DATE = sys.argv[1] if len(sys.argv) > 1 else dt.date.today().isoformat()
COMMIT = sys.argv[2] if len(sys.argv) > 2 else default_commit()

doc = yaml.safe_load((ROOT / "research/spec/rules.yaml").read_text())
SPEC_VERSION = doc["spec_version"]
RULES = doc["rules"]
ORDER = ["TASI-W1", "TASI-W2", "TASI-W3"]
assert list(RULES) == ORDER, f"rule set/order changed: {list(RULES)}"

OP_NAME = {"ge": "ge (≥)", "gt": "gt (>)", "le": "le (≤)", "lt": "lt (<)",
           "in_co": "in_co — [lo, hi) closed-open", "in_oo": "in_oo — (lo, hi) open-open",
           "in_cc": "in_cc — [lo, hi] closed-closed"}
BRACKETS = {"in_co": ("[", ")"), "in_oo": ("(", ")"), "in_cc": ("[", "]")}
GLOSS = {("ext60", "gt", 0): "close > EMA60", ("ema21gap", "ge", 0): "close ≥ EMA21 (reclaim)"}


def fmt_num(v):
    return str(int(v)) if float(v) == int(v) else str(v)


def resolve(rule, v):
    """One threshold element -> display string; returns (text, is_structural)."""
    if isinstance(v, str) and v.startswith("@"):
        p = v[1:]
        return f"{fmt_num(RULES[rule]['params'][p])} ← @{p} (tunable)", False
    return f"{fmt_num(v)} (structural)", True


# ---- Section 1: params (first-seen order across W1, W2, W3) ----
param_order = []
for r in ORDER:
    for p in RULES[r]["params"]:
        if p not in param_order:
            param_order.append(p)

sec1_rows = [("*title*", *[RULES[r]["title"] for r in ORDER]),
             ("*funnel steps*", *[str(len(RULES[r]["funnel"])) for r in ORDER])]
for p in param_order:
    sec1_rows.append((f"`{p}`", *[fmt_num(RULES[r]["params"][p]) if p in RULES[r]["params"] else "—"
                                  for r in ORDER]))

# ---- structural literals (for Section 1 footer + completeness) ----
# structural cell count = every literal (non-@) threshold element in any funnel
struct_cells = 0
struct_simple = {}   # (feature, op, literal) -> set(rules)   for scalar-literal predicates
struct_interval = []  # (rule, feature, op, lo_raw, hi_raw) with a literal element
for r in ORDER:
    for step in RULES[r]["funnel"]:
        v = step["value"]
        if isinstance(v, list):
            lits = [x for x in v if not (isinstance(x, str) and x.startswith("@"))]
            struct_cells += len(lits)
            if lits:
                struct_interval.append((r, step["feature"], step["op"], v[0], v[1]))
        elif not (isinstance(v, str) and v.startswith("@")):
            struct_cells += 1
            struct_simple.setdefault((step["feature"], step["op"], v), set()).add(r)

sec1_struct = []
for (feat, op, lit), rules_in in sorted(struct_simple.items()):
    gloss = GLOSS.get((feat, op, lit), "")
    label = f"**{feat} ({OP_NAME[op].split(' ')[1].strip('()')}) {fmt_num(lit)}" + (f"  =  {gloss}**" if gloss else "**")
    sec1_struct.append((label, *["✓ gate" if r in rules_in else "—" for r in ORDER]))
for (r, feat, op, lo, hi) in struct_interval:
    lo_lit = not (isinstance(lo, str) and lo.startswith("@"))
    gloss = GLOSS.get((feat, "ge", lo if lo_lit else None), "")
    label = f"**{feat} ≥ {fmt_num(lo)}  =  {gloss}**" if gloss else f"**{feat} {op} {lo},{hi}**"
    cap = ""
    if isinstance(hi, str) and hi.startswith("@"):
        p = hi[1:]
        cap = f" (upper cap {fmt_num(RULES[r]['params'][p])} ← @{p})"
    sec1_struct.append((label, *[f"✓ gate{cap}" if rr == r else "—" for rr in ORDER]))

# ---- Section 2: funnel predicates ----
sec2_rows = []
for r in ORDER:
    for i, step in enumerate(RULES[r]["funnel"], 1):
        v, op = step["value"], step["op"]
        if isinstance(v, list):
            lb, rb = BRACKETS[op]
            parts = [resolve(r, x)[0] for x in v]
            thr = f"{lb}{parts[0]}, {parts[1]}{rb}"
        else:
            thr = resolve(r, v)[0]
        sec2_rows.append((r, str(i), step["feature"], OP_NAME[op], thr,
                          step["null_action"], step.get("enabled_when", "—")))

# ---- Section 3: tags/notes ----
sec3_rows = []
for r in ORDER:
    for tag, expr in RULES[r].get("tag", {}).items():
        sec3_rows.append((r, tag, f"`{expr}`"))
    if RULES[r].get("notes"):
        sec3_rows.append((r, "*notes*", RULES[r]["notes"]))

# ---- completeness checks (generator-proven) ----
checks = []
ok = True
for r in ORDER:
    n = sum(1 for row in sec2_rows if row[0] == r)
    good = n == len(RULES[r]["funnel"])
    ok &= good
    checks.append(f"- {r}: Section-2 rows {n} == funnel length {len(RULES[r]['funnel'])} — {'PASS' if good else 'FAIL'}")
keys = [(row[0], row[1]) for row in sec2_rows]
uniq = len(keys) == len(set(keys))
ok &= uniq
checks.append(f"- every predicate appears exactly once (unique (rule, step) keys): {'PASS' if uniq else 'FAIL'}")
cells = sum(1 for row in sec1_struct for c in row[1:] if str(c).startswith("✓"))
good = cells == struct_cells
ok &= good
checks.append(f"- Section-1 structural-gate cells {cells} == structural literals in funnels {struct_cells}: {'PASS' if good else 'FAIL'}")
# params completeness: every canon param appears in Section 1
missing = [p for r in ORDER for p in RULES[r]["params"] if p not in param_order]
ok &= not missing
checks.append(f"- Section-1 params cover every canon param ({len(param_order)} keys, 0 missing): {'PASS' if not missing else 'FAIL: ' + str(missing)}")
# tags completeness
n_tags = sum(len(RULES[r].get("tag", {})) + (1 if RULES[r].get("notes") else 0) for r in ORDER)
good = n_tags == len(sec3_rows)
ok &= good
checks.append(f"- Section-3 rows {len(sec3_rows)} == canon tags+notes {n_tags}: {'PASS' if good else 'FAIL'}")
assert ok, "completeness checks FAILED:\n" + "\n".join(checks)

BANNER = (f"Generated from rules.yaml ({DATE}; spec_version {SPEC_VERSION}, commit {COMMIT}). "
          "Section 1 = tunable parameters PLUS the structural (literal) gates in plain language. "
          "Section 2 = complete funnel predicates, including structural literals. Section 3 = "
          "non-gating display tags/notes defined by the Canon. Nothing defined by the Canon is "
          "intentionally omitted. Scope boundary stated by the Canon itself: universe membership "
          "(TADAWUL main board, type==stock, exclude 9xxx + REIT/Fund/ETF/Sukuk, age>=min_years) "
          "is resolved outside rules.yaml and applied before any rule. Do not edit; regenerate instead.")

# ---- markdown ----
md = ["# TASI-W1 / TASI-W2 / TASI-W3 — canonical rule definition (complete view)", "", BANNER, "",
      "## Section 1 — Parameters (tunable) + Structural gates", "",
      "| Param | TASI-W1 | TASI-W2 | TASI-W3 |", "|---|---|---|---|"]
md += [f"| {' | '.join(row)} |" for row in sec1_rows]
md.append("| **Structural gates (not tunable)** |  |  |  |")
md += [f"| {' | '.join(row)} |" for row in sec1_struct]
md += ["", "## Section 2 — Funnel predicates (the executable definition, canonical order)", "",
       "| Rule | # | Feature | Operator | Threshold (resolved) | Null action | Enabled when |",
       "|---|---|---|---|---|---|---|"]
md += [f"| {' | '.join(row)} |" for row in sec2_rows]
md += ["", "## Section 3 — Display tags & notes (non-gating)", "",
       "| Rule | Tag | Expression / note |", "|---|---|---|"]
md += [f"| {' | '.join(row)} |" for row in sec3_rows]
md += ["", "## Completeness checks (generator-proven against the Canon)", ""] + checks
out_md = ROOT / ".claude/outputs/tasi-rules-comparison.md"
out_md.write_text("\n".join(md) + "\n")

# ---- xlsx (formatting only; every value above comes from rules.yaml) ----
TITLE = "TASI-W1 / TASI-W2 / TASI-W3 — canonical rule definition (complete view)"
F_TITLE = Font(bold=True, size=14, color="1F3864")
F_BANNER = Font(italic=True, size=9, color="595959")
F_HEAD = Font(bold=True, color="FFFFFF")
F_INFO = Font(italic=True, color="595959")
F_DASH = Font(color="BFBFBF")
F_SECTION = Font(bold=True, color="1F3864")
F_PASS = Font(bold=True, color="006100")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_BAND = {"TASI-W1": PatternFill("solid", fgColor="FFFFFF"),
             "TASI-W2": PatternFill("solid", fgColor="EAF1FB"),
             "TASI-W3": PatternFill("solid", fgColor="F2F2F2")}
FILL_PASS = PatternFill("solid", fgColor="C6EFCE")
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def clean(c):
    return str(c).replace("**", "").replace("`", "")


def sheet_top(ws, ncols, widths, with_banner):
    """Title (+ optional banner) merged across, then a blank spacer row."""
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.cell(row=1, column=1, value=TITLE).font = F_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    if with_banner:
        b = ws.cell(row=2, column=1, value=BANNER)
        b.font, b.alignment = F_BANNER, LEFT
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        ws.row_dimensions[2].height = 88
    return (4 if with_banner else 3)  # first content row (one blank spacer row)


def header_row(ws, r, labels):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=r, column=i, value=lab)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BOX
    ws.freeze_panes = ws.cell(row=r + 1, column=2)


def put(ws, r, i, val, align, band=None):
    c = ws.cell(row=r, column=i, value=clean(val))
    c.alignment, c.border = align, BOX
    if clean(val) == "—":
        c.font = F_DASH
    if band:
        c.fill = band
    return c


wb = Workbook()

# Sheet 1 — Params + Structural
ws = wb.active
ws.title = "Params+Structural"
r = sheet_top(ws, 4, [34, 26, 26, 26], with_banner=True)
header_row(ws, r, ["Param"] + ORDER)
r += 1
for row in sec1_rows:
    put(ws, r, 1, row[0], LEFT).font = F_INFO if row[0].startswith("*") else Font(bold=False)
    for i, v in enumerate(row[1:], 2):
        put(ws, r, i, v, LEFT if row[0].startswith("*") else CENTER)
    if row[0] == "*funnel steps*":
        for i in range(1, 5):
            ws.cell(row=r, column=i).font = F_INFO
    r += 1
sec = ws.cell(row=r, column=1, value="Structural gates (not tunable)")
sec.font, sec.fill = F_SECTION, FILL_SECTION
for i in range(2, 5):
    ws.cell(row=r, column=i).fill = FILL_SECTION
    ws.cell(row=r, column=i).border = BOX
sec.border = BOX
r += 1
for row in sec1_struct:
    put(ws, r, 1, row[0], LEFT).font = Font(bold=True)
    for i, v in enumerate(row[1:], 2):
        c = put(ws, r, i, v, CENTER)
        if clean(v).startswith("✓"):
            c.font = Font(bold=True, color="1F6B3A")
    r += 1

# Sheet 2 — Funnel (banded per rule, autofilter)
ws2 = wb.create_sheet("Funnel")
r = sheet_top(ws2, 7, [11, 5, 12, 30, 44, 11, 14], with_banner=False)
header_row(ws2, r, ["Rule", "#", "Feature", "Operator", "Threshold (resolved)", "Null action", "Enabled when"])
ws2.freeze_panes = ws2.cell(row=r + 1, column=1)
ws2.auto_filter.ref = f"A{r}:G{r + len(sec2_rows)}"
r += 1
for row in sec2_rows:
    band = FILL_BAND[row[0]]
    for i, v in enumerate(row, 1):
        put(ws2, r, i, v, CENTER if i in (1, 2, 6) else LEFT, band)
    r += 1

# Sheet 3 — Tags & notes
ws3 = wb.create_sheet("Tags")
r = sheet_top(ws3, 3, [11, 14, 70], with_banner=False)
header_row(ws3, r, ["Rule", "Tag", "Expression / note"])
r += 1
for row in sec3_rows:
    band = FILL_BAND[row[0]]
    for i, v in enumerate(row, 1):
        put(ws3, r, i, v, CENTER if i == 1 else LEFT, band)
    r += 1

# Sheet 4 — Completeness (generator-proven) + the scope banner again
ws4 = wb.create_sheet("Completeness")
r = sheet_top(ws4, 2, [86, 10], with_banner=True)
header_row(ws4, r, ["Check (generator-proven against the Canon)", "Result"])
r += 1
for c in checks:
    txt = c.lstrip("- ")
    body, _, verdict = txt.rpartition(" — ") if " — " in txt else txt.rpartition(": ")
    put(ws4, r, 1, body, LEFT)
    v = put(ws4, r, 2, verdict, CENTER)
    if verdict == "PASS":
        v.font, v.fill = F_PASS, FILL_PASS
    r += 1

wb.save(ROOT / ".claude/outputs/tasi-rules-comparison.xlsx")

print("\n".join(checks))
print(f"OK: wrote {out_md} and .xlsx  (spec_version {SPEC_VERSION}, {DATE}, commit {COMMIT})")
